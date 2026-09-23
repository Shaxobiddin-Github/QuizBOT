"""JSON va Word (.docx) fayllardan savol import qilish."""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field

MIN_OPTIONS = 2
MAX_OPTIONS = 10
LETTERS = "ABCDEFGHIJ"


class ImportError_(Exception):
    """Foydalanuvchiga ko'rsatiladigan import xatosi."""


@dataclass
class ImportResult:
    questions: list[dict] = field(default_factory=list)
    skipped: int = 0
    notes: list[str] = field(default_factory=list)

    @property
    def count(self) -> int:
        return len(self.questions)


# --------------------------------------------------------------------- umumiy
def _clean(text: str) -> str:
    text = text.replace(" ", " ").replace("​", "")
    return re.sub(r"\s+", " ", text).strip()


def _normalize(text: str, options: list[str], correct: int,
               explanation: str = "") -> dict | None:
    text = _clean(text)
    opts, seen = [], set()
    correct_value = options[correct] if 0 <= correct < len(options) else None
    for o in options:
        o = _clean(o)
        if not o or o.lower() in seen:
            continue
        seen.add(o.lower())
        opts.append(o)
    if not text or len(opts) < MIN_OPTIONS:
        return None
    opts = opts[:MAX_OPTIONS]
    try:
        idx = opts.index(_clean(correct_value)) if correct_value else 0
    except ValueError:
        return None
    return {
        "text": text[:1000],
        "options": opts,
        "correct": idx,
        "explanation": _clean(explanation)[:500],
    }


# ----------------------------------------------------------------------- JSON
def parse_json(raw: bytes | str) -> ImportResult:
    if isinstance(raw, bytes):
        try:
            raw = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            raw = raw.decode("cp1251", errors="replace")
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ImportError_(f"JSON o'qib bo'lmadi: {exc.msg} (satr {exc.lineno})") from exc

    if isinstance(data, dict):
        for key in ("questions", "items", "data", "tests", "savollar"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            data = [data]
    if not isinstance(data, list):
        raise ImportError_("JSON ichida savollar ro'yxati topilmadi.")

    res = ImportResult()
    for item in data:
        q = _json_item(item)
        if q:
            res.questions.append(q)
        else:
            res.skipped += 1
    if not res.questions:
        raise ImportError_("Fayldan birorta ham to'g'ri savol ajratib olinmadi.")
    return res


def _json_item(item) -> dict | None:
    if not isinstance(item, dict):
        return None
    lower = {str(k).lower(): v for k, v in item.items()}

    text = _first(lower, ["q", "question", "text", "savol", "title", "name"])
    if not text:
        return None
    explanation = _first(lower, ["explanation", "izoh", "comment", "note"]) or ""

    # 1-shakl: {"q":..., "c": to'g'ri, "a": [xatolar]}  (artifact shakli)
    correct_val = _first(lower, ["c", "correct_answer", "answer_text", "to'g'ri", "togri"])
    wrong = _first(lower, ["a", "wrong", "incorrect", "distractors", "xato"])
    if correct_val is not None and isinstance(wrong, list):
        options = [str(correct_val)] + [str(w) for w in wrong]
        return _normalize(str(text), options, 0, str(explanation))

    # 2-shakl: {"question":..., "options":[...], "answer": index yoki matn}
    options = _first(lower, ["options", "variants", "answers", "choices", "variantlar"])
    if isinstance(options, list) and options:
        if options and isinstance(options[0], dict):
            texts, correct = [], 0
            for i, o in enumerate(options):
                ol = {str(k).lower(): v for k, v in o.items()}
                texts.append(str(_first(ol, ["text", "title", "value", "option"]) or ""))
                if ol.get("correct") or ol.get("is_correct") or ol.get("right"):
                    correct = i
            return _normalize(str(text), texts, correct, str(explanation))

        options = [str(o) for o in options]
        ans = _first(lower, ["answer", "correct", "correct_option_id", "correct_index",
                             "javob", "right"])
        correct = 0
        if isinstance(ans, bool):
            correct = 0
        elif isinstance(ans, int):
            correct = ans - 1 if ans > 0 and ans > len(options) - 1 else ans
        elif isinstance(ans, str):
            s = ans.strip()
            if len(s) == 1 and s.upper() in LETTERS:
                correct = LETTERS.index(s.upper())
            elif s.isdigit():
                correct = int(s) - 1 if int(s) >= len(options) else int(s)
            else:
                matches = [i for i, o in enumerate(options)
                           if _clean(o).lower() == _clean(s).lower()]
                correct = matches[0] if matches else 0
        correct = max(0, min(correct, len(options) - 1))
        return _normalize(str(text), options, correct, str(explanation))
    return None


def _first(d: dict, keys: list[str]):
    for k in keys:
        if k in d and d[k] not in (None, ""):
            return d[k]
    return None


# ----------------------------------------------------------------------- DOCX
Q_MARK = re.compile(r"^(?:#+\s*|\d{1,4}\s*[\.\)]\s+|\d{1,4}\s*[-–]\s*savol\s*[\.\):]?\s*)", re.I)
OPT_PLUS = re.compile(r"^\s*([+\-*])\s*")
OPT_LETTER = re.compile(r"^\s*([A-Da-dА-Га-г])\s*[\)\.\-:]\s+")
OPT_PLUS_LETTER = re.compile(r"^\s*\+\s*([A-Da-d])\s*[\)\.\-:]\s*")
ANS_KEY = re.compile(
    r"^\s*(?:to'?g'?ri\s+javob|javob|answer|ответ)\s*[:\-–]\s*([A-Da-d0-9])\b", re.I)
INLINE_KEY = re.compile(r"\((?:to'?g'?ri\s+javob|javob)\s*[:\-–]?\s*([A-Da-d0-9])\)", re.I)


def docx_lines(raw: bytes) -> list[str]:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ImportError_("Serverda python-docx o'rnatilmagan.") from exc
    try:
        document = docx.Document(io.BytesIO(raw))
    except Exception as exc:
        raise ImportError_(
            "Word fayl o'qilmadi. Faqat .docx (eski .doc emas) qo'llab-quvvatlanadi.") from exc

    lines: list[str] = []
    for para in document.paragraphs:
        for piece in para.text.splitlines():
            if piece.strip():
                lines.append(piece.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            uniq = []
            for c in cells:
                if c and (not uniq or uniq[-1] != c):
                    uniq.append(c)
            if len(uniq) >= 2:
                lines.append(" ==== ".join(uniq))
            elif uniq:
                lines.append(uniq[0])
    return lines


def parse_docx(raw: bytes) -> ImportResult:
    lines = docx_lines(raw)
    if not lines:
        raise ImportError_("Word fayl bo'sh ko'rinadi.")
    if sum(1 for ln in lines if "====" in ln) >= 2:
        return _parse_delimited(lines)
    return _parse_blocks(lines)


def parse_txt(raw: bytes) -> ImportResult:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1251", errors="replace")
    lines = [ln.strip() for ln in text.splitlines()]
    if sum(1 for ln in lines if "====" in ln) >= 2:
        return _parse_delimited([ln for ln in lines if ln])
    return _parse_blocks(lines)


def _parse_delimited(lines: list[str]) -> ImportResult:
    """`Savol ==== +to'g'ri ==== xato ==== xato` ko'rinishi."""
    res = ImportResult()
    res.notes.append("Format: «====» ajratgichli (birinchi variant — to'g'ri javob).")
    for line in lines:
        if "====" not in line:
            continue
        parts = [p.strip() for p in line.split("====") if p.strip()]
        if len(parts) < 3:
            res.skipped += 1
            continue
        text, options = parts[0], parts[1:]
        correct = 0
        cleaned = []
        for i, o in enumerate(options):
            if o.startswith("+"):
                correct = i
                o = o[1:].strip()
            elif o.startswith("-"):
                o = o[1:].strip()
            cleaned.append(o)
        text = Q_MARK.sub("", text)
        q = _normalize(text, cleaned, correct)
        if q:
            res.questions.append(q)
        else:
            res.skipped += 1
    if not res.questions:
        raise ImportError_("«====» formatidagi savollar topilmadi.")
    return res


def _is_continuation(cur: dict) -> bool:
    """Markersiz satr savolning davomimi yoki variantmi?"""
    if cur["options"] or cur.get("cont", 0) >= 2:
        return False
    text = cur["text"].rstrip()
    if len(text) >= 400:
        return False
    return not text.endswith(("?", ":", ".", "!", "…", ")", ";"))


def _parse_blocks(lines: list[str]) -> ImportResult:
    """Savol + variantlar bloklari."""
    blocks: list[dict] = []
    cur: dict | None = None
    saw_plus = saw_key = saw_letter = False

    def flush():
        nonlocal cur
        if cur and len(cur["options"]) >= MIN_OPTIONS:
            blocks.append(cur)
        cur = None

    for raw in lines:
        line = raw.strip()
        if not line:
            continue

        key = ANS_KEY.match(line)
        if key and cur:
            saw_key = True
            cur["key"] = key.group(1)
            continue

        is_question = bool(Q_MARK.match(line))
        m_plus_letter = OPT_PLUS_LETTER.match(line)
        m_plus = OPT_PLUS.match(line)
        m_letter = OPT_LETTER.match(line)

        if is_question and not (m_plus_letter or (m_plus and m_plus.group(1) != "-")):
            flush()
            text = Q_MARK.sub("", line)
            inline = INLINE_KEY.search(text)
            cur = {"text": INLINE_KEY.sub("", text).strip(), "options": [], "marked": None,
                   "key": inline.group(1) if inline else None, "cont": 0}
            continue

        if cur is None:
            # Marker yo'q — birinchi satrni savol deb olamiz
            cur = {"text": line, "options": [], "marked": None, "key": None, "cont": 0}
            continue

        if m_plus_letter:
            saw_plus = saw_letter = True
            cur["marked"] = len(cur["options"])
            cur["options"].append(OPT_PLUS_LETTER.sub("", line).strip())
            continue
        if m_plus:
            sign = m_plus.group(1)
            body = OPT_PLUS.sub("", line).strip()
            if not body:
                continue
            if sign in "+*":
                saw_plus = True
                cur["marked"] = len(cur["options"])
            cur["options"].append(body)
            continue
        if m_letter:
            saw_letter = True
            cur["options"].append(OPT_LETTER.sub("", line).strip())
            continue

        # Markersiz satr: savol davomi yoki oddiy variant
        if _is_continuation(cur):
            cur["cont"] += 1
            cur["text"] = (cur["text"] + " " + line).strip()
        else:
            cur["options"].append(line)
    flush()

    res = ImportResult()
    for b in blocks:
        correct = b["marked"]
        if correct is None and b.get("key"):
            k = str(b["key"]).strip()
            correct = LETTERS.index(k.upper()) if k.upper() in LETTERS else int(k) - 1
        if correct is None:
            correct = 0
        correct = max(0, min(correct, len(b["options"]) - 1))
        q = _normalize(b["text"], b["options"], correct)
        if q:
            res.questions.append(q)
        else:
            res.skipped += 1

    if not res.questions:
        raise ImportError_(
            "Savollar ajratib olinmadi. Har bir savol alohida satrda, "
            "variantlar esa «A) …» yoki «+ / -» belgisi bilan yozilgan bo'lsin."
        )
    if saw_plus:
        res.notes.append("Format: «+» bilan belgilangan variant — to'g'ri javob.")
    elif saw_key:
        res.notes.append("Format: «Javob: B» satridan to'g'ri javob olindi.")
    elif saw_letter:
        res.notes.append("⚠️ To'g'ri javob belgisi topilmadi — birinchi variant to'g'ri deb olindi.")
    else:
        res.notes.append("⚠️ Belgisiz format — har blokning birinchi varianti to'g'ri deb olindi.")
    return res


# ------------------------------------------------------------------ dispatcher
def parse_file(filename: str, raw: bytes) -> ImportResult:
    name = (filename or "").lower()
    if name.endswith(".json"):
        return parse_json(raw)
    if name.endswith(".docx"):
        return parse_docx(raw)
    if name.endswith((".txt", ".md")):
        return parse_txt(raw)
    if name.endswith(".doc"):
        raise ImportError_(
            "Eski .doc formati qo'llab-quvvatlanmaydi.\n"
            "Word'da «Save as → .docx» qilib qayta yuboring.")
    raise ImportError_("Faqat .json, .docx yoki .txt fayllar qabul qilinadi.")


def to_json_export(title: str, questions: list[dict]) -> str:
    return json.dumps(
        {"title": title,
         "questions": [
             {"question": q["text"], "options": q["options"],
              "answer": q["correct"], "explanation": q.get("explanation", "")}
             for q in questions]},
        ensure_ascii=False, indent=2)
