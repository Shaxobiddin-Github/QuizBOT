"""To'liq integratsion test: haqiqiy handler'lar, soxta Telegram API."""
import asyncio, io, os, re, sys, json, traceback
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import harness as H
from aiogram import types

import db, seed
from handlers import ROUTERS
import handlers.classic as classic
import handlers.group as group

OK, FAIL = [], []
def check(name, cond, extra=""):
    (OK if cond else FAIL).append(name)
    print(("  ✅ " if cond else "  ❌ ") + name + (f"  → {extra}" if extra and not cond else ""))

def last_texts(s, n=1, kind="SendMessage"):
    out = [m.text for k, m, _ in s.calls if k == kind and getattr(m, "text", None)]
    return out[-n:]

def btns(s, kind="SendMessage"):
    for k, m, _ in reversed(s.calls):
        if k in (kind, "EditMessageText") and getattr(m, "reply_markup", None):
            mk = m.reply_markup
            if isinstance(mk, types.InlineKeyboardMarkup):
                return [b.callback_data for row in mk.inline_keyboard for b in row]
    return []


async def main():
    await db.connect(); await seed.ensure_default(); await seed.ensure_iq()
    bot, dp, S = H.make(ROUTERS)
    ali, vali, sardor = H.user(101, "Ali"), H.user(102, "Vali"), H.user(103, "Sardor")
    priv, grp = H.chat(101), H.chat(-1001, "supergroup", "Test guruh")

    print("\n━━━ 1. SHAXSIY CHAT: /start, menyu ━━━")
    await dp.feed_update(bot, H.upd_message("/start", ali, priv))
    t = last_texts(S, 2)
    check("/start salomlashadi", any("Assalomu alaykum" in x for x in t))
    check("guruhga qo'shish taklifi bor", any("guruhga qo'shing" in x for x in t))
    await dp.feed_update(bot, H.upd_message("❓ Yordam", ali, priv))
    check("yordam matni", "Qo'llanma" in last_texts(S)[-1])

    print("\n━━━ 2. PRO REJIM (shaxsiy) ━━━")
    await dp.feed_update(bot, H.upd_message("🧠 Pro rejim", ali, priv))
    check("setup kartasi", "Pro rejim" in last_texts(S)[-1])
    check("boshlash tugmasi", "su:pro:go" in btns(S))
    # savollar sonini 3 ga qo'yamiz
    await db.set_pref(101, "count", 3)
    await dp.feed_update(bot, H.upd_call("su:pro:go", ali, priv))
    card = last_texts(S)[-1]
    check("pro savol kartasi", "1-savol" in card and "▱" in card, card[:80])
    sess = await db.active_session(101, "pro")
    sid = sess["id"]
    q0 = await db.question(sess["q_ids"][0])
    corr = sess["settings"]["perm"][0].index(q0["correct"])

    msg_id = S.calls[-1][2].message_id
    await dp.feed_update(bot, H.upd_call(f"pro:f:{sid}:0", ali, priv, msg_id))
    check("50:50 ishladi", "<s>" in S.messages[msg_id])
    await dp.feed_update(bot, H.upd_call(f"pro:l:{sid}:0", ali, priv, msg_id))
    check("«o'rgandim» belgilandi", await db.is_learned(101, q0["id"]))
    await dp.feed_update(bot, H.upd_call(f"pro:l:{sid}:0", ali, priv, msg_id))
    check("«o'rgandim» qaytarildi", not await db.is_learned(101, q0["id"]))
    await dp.feed_update(bot, H.upd_call(f"pro:map:{sid}:0", ali, priv, msg_id))
    check("savollar xaritasi", "xaritasi" in S.messages[msg_id])
    await dp.feed_update(bot, H.upd_call(f"pro:g:{sid}:1", ali, priv, msg_id))
    check("navigatsiya (2-savolga)", "2-savol" in S.messages[msg_id])

    # 3 ta savolga javob (1- to'g'ri, qolgani xato)
    await dp.feed_update(bot, H.upd_call(f"pro:a:{sid}:0:{corr}", ali, priv, msg_id))
    rows = await db.session_answers(sid, 101)
    check("javob bazaga yozildi", len(rows) == 1 and rows[0]["is_correct"] == 1)
    edits = [m.text for k, m, _ in S.calls if k == "EditMessageText"]
    check("javobdan keyin javob ochiladi",
          any("✅ To'g'ri!" in e for e in edits), edits[-1][:80] if edits else "")
    for i in (1, 2):
        c = sess["settings"]["perm"][i]
        qq = await db.question(sess["q_ids"][i])
        wrong = (c.index(qq["correct"]) + 1) % len(c)
        await dp.feed_update(bot, H.upd_call(f"pro:a:{sid}:{i}:{wrong}", ali, priv, msg_id))
    final = S.messages[msg_id]
    check("test yakunlandi", "yakunlandi" in final, final[:120])
    check("natijada 1/3", "1 / 3" in final, final[:200])
    check("xatolar tugmasi", any("pro:rev" in b for b in btns(S)))
    await dp.feed_update(bot, H.upd_call(f"pro:rev:{sid}:w:0", ali, priv, msg_id))
    check("xatolar tahlili", "Xatolar" in S.messages[msg_id] and "To'g'ri:" in S.messages[msg_id])
    await dp.feed_update(bot, H.upd_call(f"pro:rev:{sid}:all:0", ali, priv, msg_id))
    check("hamma javoblar ro'yxati", "Barcha javoblar" in S.messages[msg_id])
    s2 = await db.active_session(101, "pro")
    check("sessiya yopildi", s2 is None)

    print("\n━━━ 3. XATOLAR USTIDA ISHLASH ━━━")
    await dp.feed_update(bot, H.upd_call("su:pro:mistakes", ali, priv))
    check("xato savollardan test", "Xatolar ustida ish" in last_texts(S)[-1])
    ms = await db.active_session(101, "pro")
    check("faqat xato savollar olindi", ms and len(ms["q_ids"]) == 2, str(ms and ms['q_ids']))
    await dp.feed_update(bot, H.upd_call(f"pro:fin:{ms['id']}", ali, priv,
                                         S.calls[-1][2].message_id))
    await db.abort_active(101)

    print("\n━━━ 4. KLASSIK REJIM (shaxsiy, quiz-poll) ━━━")
    await db.set_pref(101, "timer", 10)
    await dp.feed_update(bot, H.upd_message("/quiz", ali, priv))
    check("klassik setup", "Klassik rejim" in last_texts(S)[-1])
    await dp.feed_update(bot, H.upd_call("su:classic:go", ali, priv))
    await asyncio.sleep(2.0)
    check("poll yuborildi", len(S.polls) >= 1, f"{len(S.polls)} ta")
    pid = list(S.polls)[-1]
    p = S.polls[pid]
    check("quiz tipida", p.type == "quiz" and p.correct_option_id is not None)
    check("poll savol limiti", len(p.question) <= 300)
    check("poll variant limiti", all(len(o) <= 100 for o in p.options))
    # 3 ta savolga to'g'ri javob beramiz
    for _ in range(3):
        await asyncio.sleep(0.4)
        pid = list(S.polls)[-1]
        await dp.feed_update(bot, H.upd_poll_answer(pid, ali, [S.polls[pid].correct_option_id]))
        await asyncio.sleep(1.6)
    await asyncio.sleep(1.5)
    res = [t for t in S.messages.values() if "Test yakunlandi" in t]
    check("klassik natija chiqdi", bool(res), str(list(S.messages.values())[-1])[:100])
    check("3/3 to'g'ri", bool(res) and "3 / 3" in res[-1], res[-1][:120] if res else "")

    print("\n━━━ 5. GURUH: lobbi + PRO JANG ━━━")
    await dp.feed_update(bot, H.upd_message("/pro", ali, grp))
    lob = last_texts(S)[-1]
    check("lobbi ochildi", "Guruh testi" in lob and "Pro jang" in lob, lob[:100])
    lob_mid = S.calls[-1][2].message_id
    for u in (ali, vali, sardor):
        await dp.feed_update(bot, H.upd_call("g:join", u, grp, lob_mid))
    # lobbi yangilanishi flood controldan qochish uchun yig'ib yuboriladi
    await asyncio.sleep(3.2)
    check("3 qatnashchi qo'shildi", "Qatnashchilar (3)" in S.messages[lob_mid],
          S.messages[lob_mid][:200])
    await dp.feed_update(bot, H.upd_call("g:cnt", vali, grp, lob_mid))
    check("begona sozlay olmaydi", "Nechta savol" not in S.messages[lob_mid])
    await dp.feed_update(bot, H.upd_call("g:cnt", ali, grp, lob_mid))
    await dp.feed_update(bot, H.upd_call("g:set:cnt:3", ali, grp, lob_mid))
    await dp.feed_update(bot, H.upd_call("g:set:tmr:10", ali, grp, lob_mid))
    check("sozlamalar saqlandi", "Savollar: <b>3</b>" in S.messages[lob_mid],
          S.messages[lob_mid][:200])
    await dp.feed_update(bot, H.upd_call("g:go", ali, grp, lob_mid))
    await asyncio.sleep(2.6)
    gs = await db.fetch_one("SELECT id FROM sessions WHERE chat_id=-1001 AND mode='battle'"
                            " ORDER BY id DESC LIMIT 1")
    gsid = gs["id"]
    check("jang sessiyasi yaratildi", gsid is not None)
    for qi in range(3):
        for _ in range(40):
            if (gsid, qi) in group.BATTLE: break
            await asyncio.sleep(0.1)
        ctx = group.BATTLE.get((gsid, qi))
        if not ctx:
            check(f"{qi+1}-savol chiqdi", False, "BATTLE ctx yo'q"); break
        check(f"{qi+1}-savol chiqdi", True)
        # Ali to'g'ri, Vali to'g'ri, Sardor xato
        await dp.feed_update(bot, H.upd_call(f"gb:{gsid}:{qi}:{ctx['correct']}", ali, grp, ctx["msg_id"]))
        await dp.feed_update(bot, H.upd_call(f"gb:{gsid}:{qi}:{ctx['correct']}", ali, grp, ctx["msg_id"]))
        await asyncio.sleep(0.05)
        await dp.feed_update(bot, H.upd_call(f"gb:{gsid}:{qi}:{ctx['correct']}", vali, grp, ctx["msg_id"]))
        bad = (ctx["correct"] + 1) % len(ctx["shown"])
        await dp.feed_update(bot, H.upd_call(f"gb:{gsid}:{qi}:{bad}", sardor, grp, ctx["msg_id"]))
        await asyncio.sleep(3.4)
    rows = await db.session_answers(gsid)
    check("takroriy javob bloklandi", len([r for r in rows if r["user_id"] == 101]) == 3,
          f"{len([r for r in rows if r['user_id']==101])} ta")
    check("3 kishidan 9 ta javob", len(rows) == 9, f"{len(rows)} ta")
    fin = [t for t in S.messages.values() if "Pro jang yakunlandi" in t]
    check("jang yakuniy reytingi", bool(fin), str(list(S.messages.values())[-1])[:120])
    if fin:
        check("Ali g'olib (tezlik boncusi)", "G'olib: <b>Ali</b>" in fin[-1], fin[-1][-200:])
    await dp.feed_update(bot, H.upd_message("/reyting", ali, grp))
    check("/reyting ishladi", "reytingi" in last_texts(S)[-1])

    # poyga holati: o'yin boshlangach «Qatnashaman» bosilsa yiqilmasin
    await dp.feed_update(bot, H.upd_call("g:join", sardor, grp, lob_mid))
    await dp.feed_update(bot, H.upd_call("g:cnt", ali, grp, lob_mid))
    await dp.feed_update(bot, H.upd_call("g:back", vali, grp, lob_mid))
    check("yopilgan lobbida yiqilmaydi", True)

    print("\n━━━ 6. GURUH: KLASSIK ━━━")
    await dp.feed_update(bot, H.upd_message("/quiz", ali, grp))
    gmid = S.calls[-1][2].message_id
    for u in (ali, vali):
        await dp.feed_update(bot, H.upd_call("g:join", u, grp, gmid))
    await dp.feed_update(bot, H.upd_call("g:set:cnt:2", ali, grp, gmid))
    await dp.feed_update(bot, H.upd_call("g:set:tmr:5", ali, grp, gmid))
    n_polls = len(S.polls)
    await dp.feed_update(bot, H.upd_call("g:go", ali, grp, gmid))
    await asyncio.sleep(3.0)
    check("guruhda poll chiqdi", len(S.polls) > n_polls)
    pid = list(S.polls)[-1]
    await dp.feed_update(bot, H.upd_poll_answer(pid, ali, [S.polls[pid].correct_option_id]))
    await dp.feed_update(bot, H.upd_poll_answer(pid, vali, [(S.polls[pid].correct_option_id + 1) % 4]))
    await asyncio.sleep(14.0)
    board = [t for t in S.messages.values() if "Reyting" in t and "🥇" in t]
    check("guruh reytingi chiqdi", bool(board), str(list(S.messages.values())[-1])[:120])
    await dp.feed_update(bot, H.upd_message("/stop", ali, grp))

    print("\n━━━ 7. IMPORT: JSON / DOCX / MATN ━━━")
    import docx
    import tempfile
    tmp = tempfile.mkdtemp()
    d = docx.Document()
    for l in ["1. Test savol bir?", "+To'g'ri", "-Xato1", "-Xato2"]:
        d.add_paragraph(l)
    d.save(f"{tmp}/imp.docx")
    jsonb = json.dumps([{"question": "JSON savol?", "options": ["a", "b", "c"], "answer": 2}])

    async def fake_download(doc, *a, **k):
        return io.BytesIO(open(f"{tmp}/imp.docx", "rb").read() if doc.file_name.endswith(".docx")
                          else jsonb.encode())
    bot.download = fake_download

    await dp.feed_update(bot, H.upd_message("➕ Savol qo'shish", ali, priv))
    t = last_texts(S)[-1]
    check("baza tanlash chiqdi", "qaysi bazaga" in t or "Sizda hali baza yo'q" in t, t[:80])
    await dp.feed_update(bot, H.upd_call("lib:new", ali, priv))
    await dp.feed_update(bot, H.upd_message("Sinov bazasi", ali, priv))
    check("yangi baza yaratildi", "bazasi yaratildi" in last_texts(S)[-1])
    newcol = (await db.get_prefs(101))["collection_id"]

    doc = types.Document(file_id="f1", file_unique_id="u1", file_name="imp.docx", file_size=1000)
    upd = H.upd_message("", ali, priv, document=doc)
    await dp.feed_update(bot, upd)
    prev = [t for t in S.messages.values() if "Tahlil natijasi" in t]
    check("docx tahlil qilindi", bool(prev) and "1</b> ta savol" in prev[-1],
          prev[-1][:150] if prev else "")
    await dp.feed_update(bot, H.upd_call("lib:save", ali, priv, S.calls[-1][1].message_id
                                         if hasattr(S.calls[-1][1], "message_id") else None))
    check("docx saqlandi", await db.count_questions(newcol) == 1,
          str(await db.count_questions(newcol)))

    doc2 = types.Document(file_id="f2", file_unique_id="u2", file_name="imp.json", file_size=100)
    await dp.feed_update(bot, H.upd_message("", ali, priv, document=doc2))
    mid2 = [m for k, m, _ in S.calls if k == "EditMessageText"][-1].message_id
    await dp.feed_update(bot, H.upd_call("lib:save", ali, priv, mid2))
    check("json saqlandi", await db.count_questions(newcol) == 2,
          str(await db.count_questions(newcol)))
    # takror
    await dp.feed_update(bot, H.upd_message("", ali, priv, document=doc2))
    mid3 = [m for k, m, _ in S.calls if k == "EditMessageText"][-1].message_id
    await dp.feed_update(bot, H.upd_call("lib:save", ali, priv, mid3))
    check("takror savol filtrlandi", await db.count_questions(newcol) == 2,
          str(await db.count_questions(newcol)))

    await dp.feed_update(bot, H.upd_call(f"lib:addto:{newcol}", ali, priv))
    await dp.feed_update(bot, H.upd_message("1. Matndan savol?\n+ha\n-yo'q\n-balki", ali, priv))
    mid4 = [m for k, m, _ in S.calls if k == "EditMessageText"][-1].message_id
    await dp.feed_update(bot, H.upd_call("lib:save", ali, priv, mid4))
    check("matndan qo'shildi", await db.count_questions(newcol) == 3,
          str(await db.count_questions(newcol)))

    print("\n━━━ 8. BAZALAR / EKSPORT / STATISTIKA ━━━")
    await dp.feed_update(bot, H.upd_message("📚 Bazalar", ali, priv))
    check("bazalar ro'yxati", "Bazalar" in last_texts(S)[-1])
    await dp.feed_update(bot, H.upd_call(f"lib:open:{newcol}", ali, priv, S.calls[-1][2].message_id))
    await dp.feed_update(bot, H.upd_call(f"lib:export:{newcol}", ali, priv))
    check("JSON eksport", any(k == "SendDocument" for k, _, _ in S.calls))
    await dp.feed_update(bot, H.upd_call(f"lib:del:{newcol}", ali, priv))
    await dp.feed_update(bot, H.upd_call(f"lib:del2:{newcol}", ali, priv))
    check("baza o'chirildi", await db.collection(newcol) is None)
    check("default baza saqlanib qoldi",
          (await db.count_questions(await db.default_collection_id())) == 197)
    await dp.feed_update(bot, H.upd_message("📊 Statistika", ali, priv))
    st = last_texts(S)[-1]
    check("statistika", "statistikangiz" in st and "Aniqlik" in st, st[:120])

    print("\n━━━ 9. CHEGARAVIY HOLATLAR ━━━")
    await dp.feed_update(bot, H.upd_call("pro:a:999999:0:0", ali, priv))
    check("yo'q sessiyada yiqilmaydi", True)
    await dp.feed_update(bot, H.upd_call(f"gb:{gsid}:0:0", ali, grp))
    check("tugagan savolda yiqilmaydi", True)
    await dp.feed_update(bot, H.upd_message("/stop", vali, priv))
    check("faol test yo'q javobi", "faol test yo'q" in last_texts(S)[-1].lower())
    bad_doc = types.Document(file_id="f3", file_unique_id="u3", file_name="x.doc", file_size=10)
    await dp.feed_update(bot, H.upd_message("", ali, priv, document=bad_doc))
    check(".doc rad etildi", any(".doc" in t and "qo'llab" in t for t in S.messages.values()))

    print("\n━━━ 10. RUXSATLAR: kim kimning savolini ko'radi ━━━")
    begona = H.user(104, "Begona")
    H.MEMBERS[-1001] = {101, 102}          # Sardor va Begona guruhda emas
    await db.touch_user(104, "Begona", None)
    import access
    access._member_cache.clear()

    # Ali yangi shaxsiy baza ochadi
    await dp.feed_update(bot, H.upd_call("lib:new", ali, priv))
    await dp.feed_update(bot, H.upd_message("Pedagogika 402", ali, priv))
    mycol = (await db.get_prefs(101))["collection_id"]
    col = await db.collection(mycol)
    check("yangi baza private", col["visibility"] == "private", col["visibility"])
    check("egasi Ali", col["owner_id"] == 101)
    await dp.feed_update(bot, H.upd_message("1. Yopiq savol?\n+ha\n-yoq\n-balki", ali, priv))
    mid5 = [m for k, m, _ in S.calls if k == "EditMessageText"][-1].message_id
    await dp.feed_update(bot, H.upd_call("lib:save", ali, priv, mid5))
    check("savol qo'shildi", await db.count_questions(mycol) == 1)

    vis_ali = {c["id"] for c in await access.visible_collections(bot, 101)}
    vis_vali = {c["id"] for c in await access.visible_collections(bot, 102)}
    check("Ali o'z bazasini ko'radi", mycol in vis_ali)
    check("Vali ko'ra olmaydi", mycol not in vis_vali, str(vis_vali))

    # guruhga ochamiz
    await dp.feed_update(bot, H.upd_call(f"lib:perm:{mycol}", ali, priv))
    pmid = [m for k, m, _ in S.calls if k == "EditMessageText"][-1].message_id
    check("ruxsatlar paneli", "kim ko'ra oladi" in S.messages[pmid].lower())
    await dp.feed_update(bot, H.upd_call(f"lib:vis:{mycol}:groups", ali, priv, pmid))
    check("guruh ro'yxati chiqdi", "Test guruh" in S.messages[pmid], S.messages[pmid][:200])
    await dp.feed_update(bot, H.upd_call(f"lib:grp:{mycol}:-1001", ali, priv, pmid))
    check("guruhga ochildi", await db.share_chats(mycol) == [-1001])

    access._member_cache.clear()
    vis_vali = {c["id"] for c in await access.visible_collections(bot, 102)}
    vis_begona = {c["id"] for c in await access.visible_collections(bot, 104)}
    check("guruhdosh Vali endi ko'radi", mycol in vis_vali)
    check("chetdagi Begona ko'rmaydi", mycol not in vis_begona, str(vis_begona))
    check("Begona kira olmaydi", not await access.can_access(bot, 104, mycol))

    # begona baza tanlashga urinsa
    await dp.feed_update(bot, H.upd_call(f"lib:pick:{mycol}", begona, H.chat(104)))
    check("begona tanlay olmaydi",
          (await db.get_prefs(104))["collection_id"] != mycol)
    # begona savol qo'sha olmaydi
    await dp.feed_update(bot, H.upd_call(f"lib:addto:{mycol}", begona, H.chat(104)))
    check("begona savol qo'sha olmaydi",
          any("faqat o'z bazangizga" in t for t in S.messages.values()))

    # guruh lobbisi faqat ochilgan bazalarni ko'rsatadi
    for_chat = {c["id"] for c in await db.collections_for_chat(-1001)}
    check("guruhda baza ko'rinadi", mycol in for_chat)
    for_other = {c["id"] for c in await db.collections_for_chat(-1002)}
    check("boshqa guruhda ko'rinmaydi", mycol not in for_other, str(for_other))

    # bot yangi guruhga qo'shilganda ro'yxatga tushadi
    newgrp = H.chat(-1002, "supergroup", "402-guruh")
    await dp.feed_update(bot, H.upd_my_chat_member(newgrp, ali))
    check("yangi guruh ro'yxatga olindi", (await db.chat(-1002)) is not None)
    chats = [c["chat_id"] for c in await db.user_chats(101)]
    check("guruh Ali ro'yxatida", -1002 in chats, str(chats))

    print("\n━━━ 11. OMMAVIY XABAR (broadcast) ━━━")
    from handlers import admin as admin_h
    check("Ali admin", await admin_h.is_admin(101))
    check("Vali admin emas", not await admin_h.is_admin(102))
    await dp.feed_update(bot, H.upd_message("/xabar", vali, H.chat(102)))
    check("begona /xabar ishlatolmaydi",
          "Ommaviy xabar" not in last_texts(S)[-1])
    await dp.feed_update(bot, H.upd_message("/xabar", ali, priv))
    check("admin menyusi", "Ommaviy xabar" in last_texts(S)[-1])
    amid = S.calls[-1][2].message_id
    await dp.feed_update(bot, H.upd_call("ad:pick:0", ali, priv, amid))
    check("guruhlar ro'yxati", "Qaysi guruhga" in S.messages[amid])
    await dp.feed_update(bot, H.upd_call("ad:t:one:-1002", ali, priv, amid))
    check("xabar so'raldi", "xabarni menga tashlang" in S.messages[amid])
    await dp.feed_update(bot, H.upd_message("Salom 402-guruh, ertaga imtihon!", ali, priv))
    check("tasdiq so'raldi", "Tasdiqlaysizmi" in last_texts(S)[-1])
    cmid = S.calls[-1][2].message_id
    n_copy = sum(1 for k, _, _ in S.calls if k == "CopyMessage")
    await dp.feed_update(bot, H.upd_call("ad:go", ali, priv, cmid))
    await asyncio.sleep(1.2)
    copies = [m for k, m, _ in S.calls if k == "CopyMessage"]
    check("xabar yuborildi", len(copies) == n_copy + 1, f"{len(copies)} ta")
    check("to'g'ri guruhga", bool(copies) and copies[-1].chat_id == -1002)
    check("hisobot chiqdi", "Yuborish tugadi" in S.messages.get(cmid, ""),
          S.messages.get(cmid, "")[:100])

    print("\n━━━ 11b. BOT FAOLIYATI (/faoliyat) ━━━")
    await dp.feed_update(bot, H.upd_message("/faoliyat", vali, H.chat(102)))
    check("begona ko'ra olmaydi", "Bot faoliyati" not in last_texts(S)[-1])
    await dp.feed_update(bot, H.upd_message("/faoliyat", ali, priv))
    over = last_texts(S)[-1]
    check("umumiy bo'lim", "Bot faoliyati" in over and "Foydalanuvchilar" in over,
          over[:80])
    check("rejimlar ko'rsatilgan", "Rejimlar" in over and "Klassik" in over, over[:200])
    fmid = S.calls[-1][2].message_id
    await dp.feed_update(bot, H.upd_call("act:grp", ali, priv, fmid))
    check("guruhlar bo'limi", "Guruhlar" in S.messages[fmid] and "a'zo" in S.messages[fmid],
          S.messages[fmid][:100])
    gkb = [m for k, m, _ in S.calls
           if k == "EditMessageText" and getattr(m, "reply_markup", None)][-1].reply_markup
    urls = [b.url for row in gkb.inline_keyboard for b in row if b.url]
    check("guruhga o'tish tugmasi", bool(urls) and "t.me" in urls[0], str(urls))
    await dp.feed_update(bot, H.upd_call("act:col", ali, priv, fmid))
    check("bazalar bo'limi", "Bazalar" in S.messages[fmid] and "savol" in S.messages[fmid])
    await dp.feed_update(bot, H.upd_call("act:top", ali, priv, fmid))
    check("top bo'limi", "faol foydalanuvchilar" in S.messages[fmid].lower(),
          S.messages[fmid][:100])
    before = S.messages[fmid]
    await dp.feed_update(bot, H.upd_call("act:grp", begona, priv, fmid))
    check("begona bo'limlarni ocha olmaydi", S.messages[fmid] == before)

    print("\n━━━ 11c. QADASH / FAYLLAR / TOZALASH ━━━")
    import recorder
    # yozib borish: guruhga yuborilgan xabarlar bazaga tushdimi
    rec = await db.fetch_one(
        "SELECT COUNT(*) n FROM sent_messages WHERE chat_id=-1001")
    check("guruh xabarlari yozib borilmoqda", rec["n"] > 0, str(rec["n"]))

    # broadcast'da qadash
    await dp.feed_update(bot, H.upd_message("/xabar", ali, priv))
    bmid = S.calls[-1][2].message_id
    await dp.feed_update(bot, H.upd_call("ad:t:one:-1001", ali, priv, bmid))
    await dp.feed_update(bot, H.upd_message("E'lon: ertaga imtihon", ali, priv))
    kmid = S.calls[-1][2].message_id
    kb = [b.callback_data for row in S.calls[-1][1].reply_markup.inline_keyboard
          for b in row]
    check("qadash tugmasi bor", "ad:pin" in kb, str(kb))
    await dp.feed_update(bot, H.upd_call("ad:pin", ali, priv, kmid))
    npin = sum(1 for k, _, _ in S.calls if k == "PinChatMessage")
    await dp.feed_update(bot, H.upd_call("ad:go", ali, priv, kmid))
    await asyncio.sleep(1.2)
    pins = [m for k, m, _ in S.calls if k == "PinChatMessage"]
    check("xabar qadaldi", len(pins) == npin + 1, f"{len(pins)} ta")
    check("to'g'ri guruhda qadaldi", bool(pins) and pins[-1].chat_id == -1001)
    pinned_row = await db.fetch_one(
        "SELECT COUNT(*) n FROM sent_messages WHERE chat_id=-1001 AND pinned=1")
    check("qadalgani bazada belgilandi", pinned_row["n"] >= 1, str(pinned_row["n"]))

    # /fayllar
    await dp.feed_update(bot, H.upd_message("/fayllar", ali, grp))
    check("fayl yo'q deb aytadi", "fayl yig'ilmagan" in last_texts(S)[-1],
          last_texts(S)[-1][:80])
    await db.record_sent(-1001, 999001, "document", "FILEID1", "maruza.pdf", "")
    # a'zo tashlagan fayl ham arxivga tushadi va tozalashdan himoyalanadi
    await db.record_sent(-1001, 999003, "document", "FILEID2", "amaliy.docx", "",
                         author="Vali", keep=1)
    await dp.feed_update(bot, H.upd_message("/fayllar", ali, grp))
    ftext = last_texts(S)[-1]
    check("fayllar ro'yxati", "maruza.pdf" in ftext, ftext[:100])
    check("a'zo fayli va muallifi", "amaliy.docx" in ftext and "Vali" in ftext,
          ftext[:200])
    check("xabarga havola", "t.me/c/" in ftext)
    ndoc = sum(1 for k, _, _ in S.calls if k == "SendDocument")
    await dp.feed_update(bot, H.upd_call("gf:-1001:999001", ali, grp))
    check("fayl qayta yuborildi",
          sum(1 for k, _, _ in S.calls if k == "SendDocument") == ndoc + 1)

    # /tozalash
    await dp.feed_update(bot, H.upd_message("/tozalash", vali, H.chat(102)))
    check("begona tozalashni ko'rmaydi", "Avtomatik tozalash" not in last_texts(S)[-1])
    await dp.feed_update(bot, H.upd_message("/tozalash", ali, priv))
    cl = last_texts(S)[-1]
    check("tozalash paneli", "Avtomatik tozalash" in cl and "soat" in cl, cl[:80])
    cmid2 = S.calls[-1][2].message_id
    await dp.feed_update(bot, H.upd_call("cl:h:12", ali, priv, cmid2))
    check("muddat o'zgardi",
          (await db.get_setting(recorder.KEY_HOURS)) == "12",
          await db.get_setting(recorder.KEY_HOURS))
    await dp.feed_update(bot, H.upd_call("cl:toggle", ali, priv, cmid2))
    check("o'chirish tugmasi ishlaydi",
          (await db.get_setting(recorder.KEY_ENABLED)) == "0")
    await dp.feed_update(bot, H.upd_call("cl:toggle", ali, priv, cmid2))
    check("qayta yoqildi", (await db.get_setting(recorder.KEY_ENABLED)) == "1")

    # haqiqiy tozalash: eski xabar o'chadi, qadalgani qoladi
    await db.execute(
        "UPDATE sent_messages SET sent_at=? WHERE chat_id=-1001", (db.iso_ago(days=5),))
    before = await db.fetch_one("SELECT COUNT(*) n FROM sent_messages WHERE chat_id=-1001")
    ndel = sum(1 for k, _, _ in S.calls if k == "DeleteMessage")
    done, failed = await recorder.clean_once(bot)
    check("eski xabarlar o'chirildi", done > 0, f"done={done}")
    check("DeleteMessage chaqirildi",
          sum(1 for k, _, _ in S.calls if k == "DeleteMessage") > ndel)
    left = await db.fetch_all(
        "SELECT pinned, keep, file_name FROM sent_messages WHERE chat_id=-1001")
    check("qadalgan/saqlangan xabarlar o'chmadi",
          bool(left) and all(r["pinned"] == 1 or r["keep"] == 1 for r in left),
          f"qolgan={[dict(r) for r in left]}")
    check("a'zo fayli tozalashdan saqlandi",
          any(r["file_name"] == "amaliy.docx" for r in left))
    after = await db.fetch_one("SELECT COUNT(*) n FROM sent_messages WHERE chat_id=-1001")
    check("ro'yxat qisqardi", after["n"] < before["n"], f"{before['n']} -> {after['n']}")

    # o'chiq holatda tozalamasin
    await db.set_setting(recorder.KEY_ENABLED, "0")
    await db.record_sent(-1001, 999002, "text", "", "", "eski")
    await db.execute("UPDATE sent_messages SET sent_at=? WHERE message_id=999002",
                     (db.iso_ago(days=5),))
    d2, _ = await recorder.clean_once(bot)
    check("o'chiq holatda tozalamaydi", d2 == 0, str(d2))
    await db.set_setting(recorder.KEY_ENABLED, "1")

    print("\n━━━ 12. IQ TEST (standart tuzilma + rasmlar) ━━━")
    from handlers import iq as iq_h
    import iq_score, media

    def last_sent():
        for k, m, r in reversed(S.calls):
            if k in ("SendMessage", "SendPhoto") and r.message_id not in H.DELETED:
                return r.message_id

    iqcol = await iq_h.iq_collection()
    n_iq = await db.count_questions(iqcol) if iqcol else 0
    check("IQ bazasi yuklandi (matn + rasm)", n_iq >= 120, str(n_iq))
    n_img = (await db.fetch_one("SELECT COUNT(*) AS n FROM questions WHERE collection_id=?"
                                " AND option_images != ''", (iqcol,)))["n"]
    check("rasmli matritsalar bazada", n_img >= 60, str(n_img))
    anyq = await db.fetch_one("SELECT id FROM questions WHERE collection_id=? AND "
                              "option_images != '' ORDER BY id LIMIT 1", (iqcol,))
    qd = await db.question(anyq["id"])
    check("rasm fayllari diskda", media.is_local(qd["image"])
          and all(media.is_local(x) for x in qd["option_images"]))
    check("rasmli variantlar aralashtirilmaydi",
          db.make_perm(qd) == list(range(len(qd["options"]))))
    await seed.ensure_iq()
    check("qayta ishga tushganda takrorlanmaydi", await db.count_questions(iqcol) == n_iq)
    quiz_vis = {c["id"] for c in await access.visible_collections(bot, 101)}
    check("IQ bazasi oddiy ro'yxatda yo'q", iqcol not in quiz_vis)

    await dp.feed_update(bot, H.upd_message("🧩 IQ test", ali, priv))
    check("IQ qoidalari", "IQ test" in last_texts(S)[-1] and "daqiqa" in last_texts(S)[-1])
    check("ogohlantirish bor", "rasmiy" in last_texts(S)[-1].lower())
    n_calls = len(S.calls)
    await dp.feed_update(bot, H.upd_call("iq:go", ali, priv))
    iqs = await db.active_session(101, "iq")
    check("IQ sessiyasi (30 savol)", iqs is not None and len(iqs["q_ids"]) == iq_h.QUESTION_COUNT,
          str(len(iqs["q_ids"])) if iqs else "yo'q")
    iqsid = iqs["id"]
    iqqs = await db.questions_by_ids(iqs["q_ids"])
    diffs = [iqqs[q]["difficulty"] for q in iqs["q_ids"]]
    check("osondan qiyinga tartiblangan", diffs == sorted(diffs), str(diffs))
    cats = [iqqs[q]["category"] for q in iqs["q_ids"]]
    check("12 ta Raven matritsasi", cats.count("Matritsalar (Raven)") == 12, str(cats))
    check("uchala soha qamralgan",
          {iq_h._domain_of(c) for c in cats} >= set(iq_h.DOMAINS), str(set(cats)))
    check("qiyinlik 1..5 qamralgan", set(diffs) >= {1, 2, 3, 4, 5}, str(set(diffs)))
    cur = last_sent()
    check("IQ savol kartasi", "IQ test" in S.messages[cur] and "⏳" in S.messages[cur],
          S.messages[cur][:100])

    # hammasiga to'g'ri javob beramiz (rasm <-> matn almashishi bilan)
    for i, qid in enumerate(iqs["q_ids"]):
        order = iqs["settings"]["perm"][i]
        corr = order.index(iqqs[qid]["correct"])
        await dp.feed_update(bot, H.upd_call(f"iq:a:{iqsid}:{i}:{corr}", ali, priv, cur))
        cur = last_sent()
    kinds = [k for k, _, _ in S.calls[n_calls:]]
    check("rasmli savollar yuborildi", "SendPhoto" in kinds)
    check("rasm → rasm tahrirlandi", "EditMessageMedia" in kinds)
    check("rasm ↔ matn xabari almashtirildi", "DeleteMessage" in kinds)
    check("file_id keshlandi",
          (await db.fetch_one("SELECT COUNT(*) AS n FROM media_cache"))["n"] > 0)
    res = S.messages[cur]
    check("IQ natija chiqdi", "IQ test yakunlandi" in res, res[:100])
    m = re.search(r"IQ: <b>(\d+)</b>", res)
    check("hammasi to'g'ri → IQ ≥ 125", bool(m) and int(m.group(1)) >= 125, res[:200])
    check("ishonch oralig'i va persentil", "ishonch oralig'i" in res and "Persentil" in res)
    check("WAIS tasnifi", "WAIS" in res)
    check("sohalar tahlili", "Sohalar bo'yicha" in res and "Vizual-fazoviy" in res)
    check("natijada ogohlantirish", "rasmiy" in res.lower())
    await dp.feed_update(bot, H.upd_call(f"iq:rev:{iqsid}:0", ali, priv, cur))
    cur = last_sent()
    check("javoblar tahlili", "Javoblar tahlili" in S.messages[cur])
    see = [b for b in btns(S) if b.startswith("iq:see:")]
    check("rasmli savolni ko'rish tugmasi", bool(see), str(btns(S)))
    if see:
        await dp.feed_update(bot, H.upd_call(see[0], ali, priv, cur))
        cur = last_sent()
        check("rasm to'g'ri javob bilan", cur in H.PHOTO_MSGS
              and "To'g'ri javob" in S.messages[cur], S.messages[cur][:120])
    await dp.feed_update(bot, H.upd_message("🧩 IQ test", ali, priv))
    check("oxirgi natija ko'rsatiladi", "Oxirgi natijangiz" in last_texts(S)[-1])

    # vaqt tugashi
    await dp.feed_update(bot, H.upd_call("iq:go", vali, H.chat(102)))
    vmid = last_sent()
    vs = await db.active_session(102, "iq")
    past = (__import__("datetime").datetime.now(__import__("datetime").timezone.utc)
            - __import__("datetime").timedelta(minutes=1)).isoformat(timespec="seconds")
    st = dict(vs["settings"]); st["deadline"] = past
    await db.execute("UPDATE sessions SET settings=? WHERE id=?",
                     (json.dumps(st, ensure_ascii=False), vs["id"]))
    await dp.feed_update(bot, H.upd_call(f"iq:a:{vs['id']}:0:0", vali, H.chat(102), vmid))
    vs2 = await db.get_session(vs["id"])
    check("vaqt tugagach yakunlanadi", vs2["status"] in ("timeout", "done"), vs2["status"])
    check("vaqt tugagach natija", "IQ test yakunlandi" in S.messages[last_sent()],
          S.messages[last_sent()][:80])
    check("javobsiz → past IQ", re.search(r"IQ: <b>(\d+)</b>", S.messages[last_sent()])
          and int(re.search(r"IQ: <b>(\d+)</b>", S.messages[last_sent()]).group(1)) < 80)

    print("\n━━━ 13. IQ HISOBI (IRT 3PL) ━━━")
    def rep_for(k):
        items = [iq_score.Item(d, 6, i < k) for i, d in
                 enumerate([1] * 6 + [2] * 6 + [3] * 6 + [4] * 6 + [5] * 6)]
        # osonlaridan boshlab k tasi to'g'ri
        return iq_score.report(items)
    scores = [rep_for(k).iq for k in (0, 6, 12, 18, 24, 30)]
    check("ball monoton o'sadi", scores == sorted(scores) and len(set(scores)) == 6, str(scores))
    mid = rep_for(15)
    check("yarmi to'g'ri → o'rtacha (90–110)", 90 <= mid.iq <= 110, str(mid.iq))
    check("ishonch oralig'i ichida", mid.ci_low <= mid.iq <= mid.ci_high)
    check("persentil 100 → 50%", abs(iq_score.percentile(0) - 50) < 1e-9)
    check("WAIS toifalari", iq_score.classify(131)[0] == "Juda yuqori"
          and iq_score.classify(100)[0] == "O'rtacha" and iq_score.classify(65)[0] == "Past")
    hard = iq_score.report([iq_score.Item(5, 8, True)] * 5 + [iq_score.Item(1, 8, False)] * 5)
    easy = iq_score.report([iq_score.Item(1, 8, True)] * 5 + [iq_score.Item(5, 8, False)] * 5)
    check("qiyin savol ko'proq ball beradi", hard.iq != easy.iq and hard.se > easy.se * 0.5)

    print("\n━━━ 14. TUZATILGAN XATOLAR ━━━")
    # a) shaxsiy test guruh o'yinini to'xtatmasligi kerak
    access._member_cache.clear()
    H.MEMBERS.pop(-1001, None)
    await dp.feed_update(bot, H.upd_message("/pro", ali, grp))
    lmid = S.calls[-1][2].message_id
    await dp.feed_update(bot, H.upd_call("g:join", ali, grp, lmid))
    await dp.feed_update(bot, H.upd_call("g:set:tmr:30", ali, grp, lmid))
    await dp.feed_update(bot, H.upd_call("g:go", ali, grp, lmid))
    await asyncio.sleep(2.3)
    g = await db.fetch_one("SELECT id FROM sessions WHERE chat_id=-1001 AND status='active'"
                           " ORDER BY id DESC LIMIT 1")
    check("guruh jangi boshlandi", g is not None)
    await dp.feed_update(bot, H.upd_call("su:pro:go", ali, priv))
    await dp.feed_update(bot, H.upd_message("/stop", ali, priv))
    gs = await db.get_session(g["id"]) if g else None
    check("shaxsiy test/stop guruh o'yiniga tegmaydi", gs and gs["status"] == "active",
          gs and gs["status"])
    await dp.feed_update(bot, H.upd_message("/stop", ali, grp))
    gs = await db.get_session(g["id"]) if g else None
    check("guruhda /stop ishlaydi", gs and gs["status"] == "aborted")

    # b) default baza o'chmaydi
    dcol = await db.default_collection_id()
    await db.execute("UPDATE collections SET owner_id=101 WHERE id=?", (dcol,))
    await dp.feed_update(bot, H.upd_call(f"lib:open:{dcol}", ali, priv))
    check("default bazada o'chirish tugmasi yo'q", f"lib:del:{dcol}" not in btns(S))
    await dp.feed_update(bot, H.upd_call(f"lib:del2:{dcol}", ali, priv))
    check("default baza savollari saqlandi", await db.count_questions(dcol) == 197,
          str(await db.count_questions(dcol)))

    # c) klassik natijadan «xatolarni ko'rish» yiqilmaydi
    cs = await db.fetch_one("SELECT id FROM sessions WHERE owner_id=101 AND chat_id=101 "
                            "AND mode='classic' ORDER BY id LIMIT 1")
    await dp.feed_update(bot, H.upd_call(f"pro:rev:{cs['id']}:all:0", ali, priv))
    check("klassik tahlil ishlaydi", "Barcha javoblar" in last_texts(S, kind="EditMessageText")[-1]
          if last_texts(S, kind="EditMessageText") else False)

    # d) begona Pro testni yakunlay olmaydi
    await dp.feed_update(bot, H.upd_call("su:pro:go", ali, priv))
    ps = await db.active_session(101, "pro")
    await dp.feed_update(bot, H.upd_call(f"pro:fin2:{ps['id']}", vali, priv))
    check("begona yakunlay olmaydi", (await db.get_session(ps["id"]))["status"] == "active")
    await dp.feed_update(bot, H.upd_message("/stop", ali, priv))
    check("POLLS tozalanadi", len(classic.POLLS) <= 1, str(len(classic.POLLS)))

    print("\n━━━ 15. RASMLI SAVOLLAR IMPORTI ━━━")
    import zipfile
    from PIL import Image, ImageDraw

    def png(color, shape="rect"):
        im = Image.new("RGB", (120, 120), "white")
        d = ImageDraw.Draw(im)
        (d.ellipse if shape == "circle" else d.rectangle)([20, 20, 100, 100], fill=color)
        b = io.BytesIO(); im.save(b, "PNG"); return b.getvalue()

    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as z:
        z.writestr("test/savollar.json", json.dumps({"questions": [
            {"question": "Qaysi rang qizil?", "option_images": ["img/a.png", "img/b.png", "img/c.png"],
             "answer": "B", "explanation": "B — qizil"},
            {"question": "Rasmdagi shakl?", "image": "img/q.png",
             "options": ["Doira", "Kvadrat", "Uchburchak"], "answer": 0},
        ]}))
        z.writestr("test/img/a.png", png("blue")); z.writestr("test/img/b.png", png("red"))
        z.writestr("test/img/c.png", png("green")); z.writestr("test/img/q.png", png("black", "circle"))
    zip_bytes = zbuf.getvalue()
    photo_bytes = png("orange", "circle")

    async def fake_download2(obj, *a, **k):
        name = getattr(obj, "file_name", None) or ""
        if name.endswith(".zip"):
            return io.BytesIO(zip_bytes)
        if isinstance(obj, types.PhotoSize) or isinstance(obj, str):
            return io.BytesIO(photo_bytes)
        return await fake_download(obj)
    bot.download = fake_download2

    await dp.feed_update(bot, H.upd_call("lib:new", ali, priv))
    await dp.feed_update(bot, H.upd_message("Rasmli baza", ali, priv))
    icol = (await db.get_prefs(101))["collection_id"]
    zdoc = types.Document(file_id="fz", file_unique_id="uz", file_name="rasmli.zip", file_size=5000)
    await dp.feed_update(bot, H.upd_message("", ali, priv, document=zdoc))
    prev = [t for t in S.messages.values() if "Tahlil natijasi" in t][-1]
    check("zip tahlil qilindi", "2</b> ta savol" in prev and "Rasmli savollar" in prev, prev[:300])
    zmid = [m for k, m, _ in S.calls if k == "EditMessageText"][-1].message_id
    await dp.feed_update(bot, H.upd_call("lib:save", ali, priv, zmid))
    check("zip saqlandi", await db.count_questions(icol) == 2, str(await db.count_questions(icol)))
    ids = await db.pick_questions(icol, 0, shuffle=False)
    qs = await db.questions_by_ids(ids)
    q1 = qs[ids[0]]
    check("variant rasmlari saqlandi", len(q1["option_images"]) == 3
          and all(media.is_local(x) for x in q1["option_images"]) and q1["correct"] == 1)
    check("savol rasmi saqlandi", media.is_local(qs[ids[1]]["image"]))

    # rasm + izoh bilan savol
    await dp.feed_update(bot, H.upd_call(f"lib:addto:{icol}", ali, priv))
    ph = [types.PhotoSize(file_id="phot1", file_unique_id="pu1", width=120, height=120)]
    await dp.feed_update(bot, H.upd_message("Bu qanday shakl?\n+Doira\n-Kvadrat\n-Romb", ali, priv,
                                            photo=ph))
    pmid2 = [m for k, m, _ in S.calls if k == "EditMessageText"][-1].message_id
    await dp.feed_update(bot, H.upd_call("lib:save", ali, priv, pmid2))
    check("rasm+izoh savoli qo'shildi", await db.count_questions(icol) == 3,
          str(await db.count_questions(icol)))

    # ichki tarmoq havolasi rad etiladi (SSRF himoyasi)
    import importers
    from handlers import library as lib_h
    bad = importers.ImportResult(questions=[{"text": "x", "options": ["a", "b"], "correct": 0,
                                             "image": "http://127.0.0.1/secret.png"}])
    check("ichki manzildan rasm olinmaydi", not await lib_h._materialize(bad, 101))

    # Pro rejimda rasmli savollar
    await db.set_pref(101, "count", 0)
    n_calls = len(S.calls)
    await dp.feed_update(bot, H.upd_call(f"lib:run:{icol}:pro", ali, priv))
    rs = await db.active_session(101, "pro")
    check("pro rasmli sessiya", rs is not None and len(rs["q_ids"]) == 3)
    cur = last_sent()
    for i, qid in enumerate(rs["q_ids"]):
        q = await db.question(qid)
        corr = rs["settings"]["perm"][i].index(q["correct"])
        await dp.feed_update(bot, H.upd_call(f"pro:a:{rs['id']}:{i}:{corr}", ali, priv, cur))
        cur = last_sent()
    kinds = [k for k, _, _ in S.calls[n_calls:]]
    check("pro: rasm yuborildi", kinds.count("SendPhoto") >= 1)
    check("pro: natija chiqdi", "3 / 3" in S.messages[cur], S.messages[cur][:120])

    # Klassik rejim: rasm + harfli poll
    await db.set_pref(101, "timer", 10)
    n_polls = len(S.polls)
    n_calls = len(S.calls)
    await dp.feed_update(bot, H.upd_call(f"lib:run:{icol}:classic", ali, priv))
    await asyncio.sleep(2.2)
    new_polls = list(S.polls.values())[n_polls:]
    check("klassik: poll chiqdi", bool(new_polls))
    kinds = [k for k, _, _ in S.calls[n_calls:]]
    check("klassik: savol rasmi poll'dan oldin", "SendPhoto" in kinds
          and kinds.index("SendPhoto") < kinds.index("SendPoll") if "SendPoll" in kinds else False)
    img_polls = [p for p in new_polls if [o if isinstance(o, str) else o.text for o in p.options]
                 == ["A", "B", "C"]]
    first = await db.question((await db.active_session(101, "classic"))["q_ids"][0])
    if media.has_option_images(first):
        check("rasmli variantlar — harfli poll", bool(img_polls))
    await dp.feed_update(bot, H.upd_message("/stop", ali, priv))
    await asyncio.sleep(0.2)

    await db.execute("DELETE FROM answers WHERE user_id IN (101,102,103)")
    await db.execute("DELETE FROM sessions WHERE owner_id IN (101,102,103)")
    await db.execute("DELETE FROM learned WHERE user_id IN (101,102,103)")
    await db.close()

    print("\n" + "═" * 46)
    print(f"  ✅ O'TDI: {len(OK)}      ❌ YIQILDI: {len(FAIL)}")
    if FAIL:
        for f in FAIL: print("   -", f)
    print("═" * 46)
    return 1 if FAIL else 0


async def _wrapped():
    try:
        return await main()
    finally:
        import contextlib
        with contextlib.suppress(Exception):
            await db.close()
        for t in asyncio.all_tasks():
            if t is not asyncio.current_task():
                t.cancel()

sys.exit(asyncio.run(_wrapped()))
